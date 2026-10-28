import openpyxl,json,collections,zipfile,pathlib
p='outputs/alex_central/Alex_Sede_Central_2026.xlsx'
w=openpyxl.load_workbook(p,data_only=True)
f=openpyxl.load_workbook(p,data_only=False)
s=json.loads(pathlib.Path('work/alex_spec.json').read_text(encoding='utf-8'))
source=json.loads(pathlib.Path('work/data.json').read_text(encoding='utf-8'))
groups=collections.defaultdict(list)
for r in source['execution_rows']:groups[(r['Proyecto'],str(int(r['Meta'])))].append(r)
checked=0
for sn,x in s['primary'].items():
 for l in x['leaves']:
  r=l['row'];sums={k:sum(t[k] for t in groups[(l['cui'],l['meta'])]) for k in ['PIM','Certificacion','Devengado']}
  for c,k in [('D','PIM'),('E','Certificacion'),('G','Devengado')]:assert abs(w[sn][f'{c}{r}'].value-sums[k])<.02,(sn,r,k)
  checked+=1
for l in s['missing']:
 r=l['row'];sums={k:sum(t[k] for t in groups[(l['cui'],l['meta'])]) for k in ['PIM','Certificacion','Devengado']}
 for c,k in [('D','PIM'),('E','Certificacion'),('G','Devengado')]:assert abs(w['Gerencias y Oficinas'][f'{c}{r}'].value-sums[k])<.02
 checked+=1
assert checked==436
pmi={r['Proyecto']:r for r in source['pmi']}
for sn,leaves in [(sn,x['leaves']) for sn,x in s['primary'].items()]+[('Gerencias y Oficinas',s['missing'])]:
 for l in leaves:
  for col,year in zip('LMNO',range(2026,2030)):
   expected=pmi.get(l['cui'],{}).get('PMI'+str(year),'Sin PMI')
   if expected in ('',None):expected='No informado'
   actual=w[sn][f"{col}{l['row']}"].value
   assert actual==expected,(sn,l['row'],year,actual,expected)
for c,v in [('E16',905026984),('F16',843143127.18),('H16',433575772.65)]:assert abs(w['ABS'][c].value-v)<.02
errors=[]
for sn,x in s['aliases'].items():
 for l in x['leaves']:
  for col in 'JKLMNOP':
   if col=='P':continue
   expected=w[l['primary_sheet']][f"{col}{l['primary_row']}"].value
   actual=w[sn][f"{col}{l['row']}"].value
   assert (actual or '')==(expected or ''),(sn,l['row'],col,actual,expected)
for sheet in w:
 for row in sheet:
  for c in row:
   if c.data_type=='e':errors.append((sheet.title,c.coordinate,c.value))
assert not errors,errors[:30]
assert all(n not in w.sheetnames for n in ['002','003','005','100','302','UEs-ABS','Hoja9','Hoja4'])
assert w['Gerencia de Infraestructura']['K10'].value=='Por completar'
assert f['Gerencia de Infraestructura']['L10'].data_type=='f'
assert str(w['SIAF']['B3'].value).startswith('2026-10-06')
with zipfile.ZipFile(p) as z:
 assert not any(n.startswith('xl/externalLinks/') for n in z.namelist())
 assert not any(n.endswith('vbaProject.bin') for n in z.namelist())
 assert len([n for n in z.namelist() if n.startswith('xl/queryTables/queryTable') and n.endswith('.xml')])==2
print(json.dumps({'metas_verificadas':checked,'proyectos':len({k[0] for k in groups}),'totales_conciliados':True,'errores_excel':0,'hojas_visibles':[t.title for t in w if t.sheet_state=='visible']},ensure_ascii=False,indent=2))
