"""Construye el reporte SINVERSION (UE 001) al 06/10/2026 sin Power Query.
Fuentes: XLS SIAF (cifras), Cartera PMI (montos 2026-2029), XLSM de Alex (agrupación por gerencia).
"""
import re, collections, warnings, pathlib, sys
import xlrd, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
warnings.filterwarnings('ignore')

ROOT = pathlib.Path(r'E:/New folder/Economia/8vo/gore/Cash data request')
REF = ROOT / 'ref docs'
OUT = ROOT / 'outputs' / 'SINVERSION_UE001_2026-10-06.xlsx'
FECHA = '06/10/2026'

# ---------------------------------------------------------------- XLS SIAF
sh = xlrd.open_workbook(REF / 'reporte20261006085201.xls').sheet_by_index(0)
HDR = re.compile(r'^(\d{4})\s+(\d{4})\s+(\d{7})\s+(\d{7})\s+(.*?)\s+(\d{2})\s+(\d{3})\s+(\d{4})\s*$')
num = lambda v: float(v) if v != '' else 0.0

MOD_RULES = [  # (texto en descripción SIAF, etiqueta)
    ('CONTRATA', 'Contrata'),
    ('ADMINISTRACION DIRECTA', 'Administración directa'),
    ('EXPEDIENTES TECNICOS', 'Expediente técnico'),
    ('PREINVERSION', 'Estudio de preinversión'),
    ('LAUDOS', 'Laudos arbitrales'),
]
def modalidad_label(code, desc):
    for k, v in MOD_RULES:
        if k in desc: return v
    if code.startswith('2.6. 3'): return 'Equipamiento'
    if re.match(r'^2\.6\. (8  1\. 4|7  1\. 6)', code): return 'Personal/bienes/servicios'
    return 'Otro (' + desc[:30] + ')'

metas = {}          # (cui, meta) -> dict
ue_total = None
cur = None; cur_ue = None
for r in range(sh.nrows):
    t = str(sh.cell_value(r, 1))
    m = re.match(r'UNIDAD EJECUTORA: (\d{3})', t)
    if m: cur_ue = m.group(1)
    h = HDR.match(t)
    if h:
        sec, prod, cui, act, nm, fn, dvf, grp = h.groups()
        key = (cui, sec)
        cur = metas.setdefault(key, dict(cui=cui, meta=sec, fn=fn, prod=nm, ue=cur_ue, pim=0, cert=0, comp=0, dev=0, pia=0, mods=collections.Counter()))
        continue
    if t.startswith('TOTAL UNIDAD EJECUTORA'):
        ue_total = {k: num(sh.cell_value(r + 1, c)) for k, c in dict(pia=6, pim=8, cert=9, comp=13, dev=17).items()}
    if re.match(r'^2\.6\.', t) and cur is not None:
        pia, pim, cert, comp, dev = (num(sh.cell_value(r, c)) for c in (5, 7, 9, 13, 17))
        cur['pia'] += pia; cur['pim'] += pim; cur['cert'] += cert; cur['comp'] += comp; cur['dev'] += dev
        lab = modalidad_label(t, str(sh.cell_value(r, 2)))
        cur['mods'][lab] += (pim if pim > 0 else 0)
        cur.setdefault('mods_any', collections.Counter())[lab] += abs(pia) + abs(pim) + abs(cert) + abs(comp) + abs(dev)
assert {m['ue'] for m in metas.values()} == {'001'}
sum_ = lambda k: sum(m[k] for m in metas.values())
for k in ('pim', 'cert', 'comp', 'dev'):
    assert abs(sum_(k) - ue_total[k]) < 0.5, (k, sum_(k), ue_total[k])

def modalidad_text(m):
    c = {k: v for k, v in m['mods'].items() if v > 0} or {k: v for k, v in m.get('mods_any', {}).items() if v > 0}
    if not c: return 'Sin datos'
    return ' / '.join(k for k, _ in sorted(c.items(), key=lambda kv: -kv[1]))

# ---------------------------------------------------------------- Cartera
cw = openpyxl.load_workbook(REF / 'Cartera_261006-090543.xlsx', data_only=True).active
tof = lambda v: float(str(v).replace(',', '')) if v not in (None, '') else None
cartera = {}
for row in cw.iter_rows(min_row=7, values_only=True):
    if row[4]:
        cartera[str(row[4]).strip()] = dict(nombre=row[7], m=[tof(row[14]), tof(row[15]), tof(row[16]), tof(row[17])], pim=tof(row[13]))

# ---------------------------------------------------------------- Alex (agrupación)
ax = openpyxl.load_workbook(REF / '0001917176_1_SINVERSION (1).xlsm', data_only=True)
CUIRE = re.compile(r'^(\d{6,7})\s*[:.]\s*(.*)')
FUNRE = re.compile(r'^(\d{2})\.\s')
placements = collections.defaultdict(list)   # cui -> [dict(sheet, path, meta, name, order)]
funcnames = {}
order = 0
def meta_of(v):
    try: return '%04d' % int(v)
    except Exception: return None

# Anexo 1: Gerencia de Infraestructura (filas 6-904)
ws = ax['Gerencia de Infraestructura']; grp = None; fun = None
for row in ws.iter_rows(min_row=5, max_row=904, max_col=3, values_only=True):
    a, b, c = row
    if isinstance(c, str) and CUIRE.match(c.strip()):
        mm = CUIRE.match(c.strip()); order += 1
        if grp and fun:
            placements[mm.group(1)].append(dict(sheet='Gerencia de Infraestructura', path=('GERENCIA DE INFRAESTRUCTURA',) + grp + (fun,), meta=meta_of(b), name=c.strip(), order=order))
        continue
    if isinstance(a, str):
        s = a.strip()
        if FUNRE.match(s): fun = s; funcnames[s[:2]] = s[4:].strip()
        elif s.startswith('001.') or s.upper().startswith('GERENCIA DE INFRA'): pass
        elif s == 'EJECUCIÓN POR ADMINISTRACIÓN DIRECTA': pass
        elif s.upper().startswith('TOTAL'): pass
        else:
            grp = (('EJECUCIÓN POR ADMINISTRACIÓN DIRECTA', s) if s.startswith('EJECUCIÓN POR ADMINISTRACIÓN DIRECTA - ') else (s,)); fun = None

# Anexo 2: Gerencias y Oficinas (UE 001: filas 6-313)
ws = ax['Gerencias y Oficinas']; ger = None; fun = None
for row in ws.iter_rows(min_row=6, max_row=313, max_col=3, values_only=True):
    a, b, c = row
    if isinstance(c, str) and CUIRE.match(c.strip()):
        mm = CUIRE.match(c.strip()); order += 1
        placements[mm.group(1)].append(dict(sheet='Gerencias y Oficinas', path=(ger,) + ((fun,) if fun else ()), meta=meta_of(b), name=c.strip(), order=order))
        continue
    if isinstance(a, str):
        s = a.strip()
        if FUNRE.match(s): fun = s; funcnames[s[:2]] = s[4:].strip()
        elif not s.upper().startswith('TOTAL'): ger = s; fun = None

def place(cui, meta):
    """Devuelve (placement|None, nota)."""
    ps = placements.get(cui)
    if not ps: return None, 'Sin clasificar: CUI no figura en el Excel de Alex (16/09)'
    ex = [p for p in ps if p['meta'] == meta and not p.get('used')]
    if ex: p = ex[0]; p['used'] = True; return p, ''
    ex = [p for p in ps if p['meta'] is None and not p.get('used')]
    if ex: p = ex[0]; p['used'] = True; return p, 'Meta sin número en Alex; asignado por CUI'
    ex = [p for p in ps if p['meta'] == meta]
    if ex: return ex[0], 'CUI/meta repetido en Alex; se usó la misma ubicación'
    return ps[0], 'Meta nueva respecto a Alex; se usó la ubicación del CUI'

# ---------------------------------------------------------------- Filas base
rows = []
for (cui, meta), m in metas.items():
    p, nota = place(cui, meta)
    car = cartera.get(cui)
    name = (p['name'] if p else (f"{cui}: {car['nombre']}" if car and car['nombre'] else f"{cui}: ({m['prod']} - nombre no disponible en fuentes)"))
    rows.append(dict(m, p=p, nota=nota, name=name, car=car))
# nombre Cartera si Alex no lo tiene (ya cubierto); ordenar
rows.sort(key=lambda r: (r['p']['order'] if r['p'] else 10**9, -r['pim'], r['cui'], r['meta']))
SIN = 'SIN CLASIFICAR (pendiente de asignar gerencia)'
for r in rows:
    if r['p'] is None:
        fname = funcnames.get(r['fn'])
        r['sheet'] = 'Sin clasificar'
        r['path'] = (SIN, f"{r['fn']}. {fname}" if fname else f"{r['fn']}. (función sin nombre en Alex)")
    else:
        r['sheet'] = r['p']['sheet']; r['path'] = r['p']['path']

# ---------------------------------------------------------------- Estilos
NAVY = PatternFill('solid', fgColor='002060'); YEL = PatternFill('solid', fgColor='FFFF00')
SAL = PatternFill('solid', fgColor='F4B084'); LBL = PatternFill('solid', fgColor='B4C6E7')
GRY = PatternFill('solid', fgColor='D9D9D9'); GRN = PatternFill('solid', fgColor='E2EFDA')
F = lambda **k: Font(name='Arial', size=k.pop('size', 9), **k)
thin = Side(style='thin', color='808080'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
NUM = '#,##0;(#,##0);"-"'; PCT = '0.0%'; MONEY = '#,##0'
STATES = ['Sin certificar', 'Solo certificado', 'Comprometido', 'Devengado total', 'Sin PIM']

wb = openpyxl.Workbook(); wb.remove(wb.active)

# ---------------------------------------------------------------- Hoja Base
base = wb.create_sheet('Base')
bh = ['Hoja', 'Gerencia / Oficina', 'Grupo (Alex)', 'Función', 'CUI', 'Meta', 'Proyecto', 'PIM 2026', 'Certificación', 'Compromiso anual', 'Devengado',
      'Saldo sin certificar (A-B)', 'Saldo sin comprometer (A-comp)', 'Saldo sin devengar (A-C)', 'Modalidad de ejecución', 'Estado situacional',
      'Monto Año 2026 (S/)', 'Monto Año 2027 (S/)', 'Monto Año 2028 (S/)', 'Monto Año 2029 (S/)', 'Base legal', 'Observación de clasificación']
base['A1'] = 'Base de datos por CUI-meta · UE 001 REGION PUNO-SEDE CENTRAL · al ' + FECHA; base['A1'].font = F(bold=True, size=12)
base['A2'] = 'Fuentes: SIAF reporte20261006085201.xls (cifras y modalidad), Cartera PMI 06/10/2026 (montos 2026-2029), Excel de Alex 16/09 (agrupación). Las hojas de anexo leen de aquí.'; base['A2'].font = F(italic=True)
BR0 = 4
for j, h in enumerate(bh, 1):
    c = base.cell(BR0, j, h); c.font = F(bold=True, color='FFFFFF'); c.fill = NAVY; c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center'); c.border = BOX
base.row_dimensions[BR0].height = 38
NOCART = 'No registrado en cartera'
for i, r in enumerate(rows):
    n = BR0 + 1 + i; r['brow'] = n
    path = r['path']
    grupo = ' > '.join(path[1:-1]) if r['sheet'] == 'Gerencia de Infraestructura' else ''
    ger = path[0]
    vals = [r['sheet'], ger, grupo, path[-1] if len(path) > 1 else '', r['cui'], int(r['meta']), r['name'], r['pim'], r['cert'], r['comp'], r['dev'],
            f'=H{n}-I{n}', f'=H{n}-J{n}', f'=H{n}-K{n}', modalidad_text(r),
            f'=IF(H{n}=0,"Sin PIM",IF(ROUND(K{n}-H{n},2)>=0,"Devengado total",IF(J{n}>0.005,"Comprometido",IF(I{n}>0.005,"Solo certificado","Sin certificar"))))']
    car = r['car']
    vals += [(car['m'][k] if car and car['m'][k] is not None else (NOCART if not car else 0)) for k in range(4)]
    vals += ['', r['nota']]
    for j, v in enumerate(vals, 1):
        c = base.cell(n, j, v); c.font = F()
        if j in (8, 9, 10, 11, 12, 13, 14, 17, 18, 19, 20): c.number_format = NUM
        if j in (8, 9, 10, 11): c.font = F(color='0000FF')
        if j == 7: c.alignment = Alignment(wrap_text=True, vertical='top')
BLAST = BR0 + len(rows)
for j, w in enumerate([22, 34, 32, 26, 9, 7, 60, 14, 14, 14, 14, 14, 14, 14, 30, 17, 14, 14, 14, 14, 12, 48], 1): base.column_dimensions[L(j)].width = w
base.freeze_panes = base.cell(BR0 + 1, 8); base.auto_filter.ref = f'A{BR0}:{L(len(bh))}{BLAST}'

# ---------------------------------------------------------------- Hojas de anexo
COLS = ['N°', 'META', 'PROYECTO', 'PIM 2026\n(A)', 'CERTIFICACIÓN\n(B)', 'SALDO\n(A-B)', 'DEVENGADO\n( C )', 'SALDO\n(A-C)', 'DEVEN\n%',
        'MODALIDAD\nDE\nEJECUCIÓN', 'ESTADO\nSITUACIONAL\nPROYECTO', 'Monto Año 2026 (S/)', 'Monto Año 2027 (S/)', 'Monto Año 2028 (S/)', 'Monto Año 2029 (S/)', 'BASE LEGAL']
WIDTHS = [6, 8, 62, 15, 15, 14, 15, 15, 9, 26, 18, 15, 15, 15, 15, 20]

class Node:
    def __init__(self, name, level): self.name = name; self.level = level; self.kids = collections.OrderedDict(); self.leaves = []; self.row = None

def build_tree(rs):
    root = Node('', 0)
    for r in rs:
        nd = root
        for lvl, nm in enumerate(r['path'], 1):
            nd = nd.kids.setdefault(nm, Node(nm, lvl))
        nd.leaves.append(r)
    return root

def write_annex(title, anexo, rs, subtitle_top=None):
    ws = wb.create_sheet(title)
    ws['C1'] = anexo; ws['C1'].font = F(bold=True, size=11)
    ws['B2'] = f'COMPORTAMIENTO DE EJECUCIÓN DE PROYECTOS DE INVERSION PUBLICA DEL AÑO FISCAL 2026 DE GERENCIAS REGIONALES Y UNIDADES EJECUTORAS · UE 001 · AL {FECHA}'; ws['B2'].font = F(bold=True)
    HR = 4
    for j, h in enumerate(COLS, 1):
        c = ws.cell(HR, j, h); c.border = BOX; c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
        c.font = F(bold=True, color='FFFFFF' if j <= 9 else '000000', size=10 if j <= 9 else 9); c.fill = NAVY if j <= 9 else LBL
        ws.column_dimensions[L(j)].width = WIDTHS[j - 1]
    ws.row_dimensions[HR].height = 42
    root = build_tree(rs)
    cur = [HR + 2]; cnt = [0]; first_leaf = HR + 2
    def sums(ref_cells, col): return '=' + '+'.join(ref_cells)
    def emit(nd):
        if nd.level > 0:
            nd.row = cur[0]; cur[0] += 1
        # hojas primero
        leaf_rows = []
        for r in nd.leaves:
            row = cur[0]; cur[0] += 1; cnt[0] += 1; leaf_rows.append(row); n = r['brow']
            ws.cell(row, 1, f'=SUBTOTAL(103,$B${first_leaf}:B{row})')
            ws.cell(row, 2, int(r['meta'])); ws.cell(row, 3, f'=Base!G{n}')
            ws.cell(row, 4, f'=Base!H{n}'); ws.cell(row, 5, f'=Base!I{n}'); ws.cell(row, 6, f'=D{row}-E{row}')
            ws.cell(row, 7, f'=Base!K{n}'); ws.cell(row, 8, f'=D{row}-G{row}'); ws.cell(row, 9, f'=IF(D{row}=0,0,G{row}/D{row})')
            ws.cell(row, 10, f'=Base!O{n}'); ws.cell(row, 11, f'=Base!P{n}')
            for k, col in enumerate('QRST'): ws.cell(row, 12 + k, f'=Base!{col}{n}')
            ws.cell(row, 16, f'=IF(Base!U{n}="","",Base!U{n})')
            for j in range(1, 17):
                c = ws.cell(row, j); c.font = F(); c.border = BOX
                if j in (4, 5, 6, 7, 8, 12, 13, 14, 15): c.number_format = NUM; c.alignment = Alignment(horizontal='right', vertical='center')
                elif j == 9: c.number_format = PCT; c.alignment = Alignment(horizontal='center', vertical='center')
                elif j in (1, 2): c.alignment = Alignment(horizontal='center', vertical='center')
                else: c.alignment = Alignment(wrap_text=True, vertical='center')
        kid_rows = [emit(k) for k in nd.kids.values()]
        if nd.level == 0: return None
        parts = {}
        for col in 'DEGH' if False else 'DEFGH':
            terms = []
            if leaf_rows: terms.append(f'SUM({col}{leaf_rows[0]}:{col}{leaf_rows[-1]})')
            terms += [f'{col}{k}' for k in kid_rows]
            ws[f'{col}{nd.row}'] = '=' + '+'.join(terms)
        ws[f'I{nd.row}'] = f'=IF(D{nd.row}=0,0,G{nd.row}/D{nd.row})'
        ws.cell(nd.row, 1, nd.name)
        fill = YEL if nd.level == 1 else (SAL if re.match(r'^\d{2}\.\s', nd.name) else GRY)
        for j in range(1, 17):
            c = ws.cell(nd.row, j); c.fill = fill; c.border = BOX; c.font = F(bold=True)
            if j in (4, 5, 6, 7, 8): c.number_format = NUM
            if j == 9: c.number_format = PCT; c.alignment = Alignment(horizontal='center')
        return nd.row
    top = [emit(k) for k in root.kids.values()]
    tr = cur[0] + 1
    ws.cell(tr, 1, f'TOTAL {title.upper()}')
    for col in 'DEFGH': ws[f'{col}{tr}'] = '=' + '+'.join(f'{col}{k}' for k in top)
    ws[f'I{tr}'] = f'=IF(D{tr}=0,0,G{tr}/D{tr})'
    for j in range(1, 17):
        c = ws.cell(tr, j); c.fill = NAVY; c.font = F(bold=True, color='FFFFFF'); c.border = BOX
        if j in (4, 5, 6, 7, 8): c.number_format = NUM
        if j == 9: c.number_format = PCT
    ws.freeze_panes = ws.cell(HR + 1, 4); ws.sheet_view.zoomScale = 90
    return ws, tr

sheets = {}
for nm, anexo in [('Gerencia de Infraestructura', 'ANEXO N° 01'), ('Gerencias y Oficinas', 'ANEXO N° 02'), ('Sin clasificar', 'ANEXO N° 03 · PROYECTOS SIN GERENCIA ASIGNADA')]:
    rs = [r for r in rows if r['sheet'] == nm]
    if rs: sheets[nm] = (write_annex(nm, anexo, rs), len(rs))

# ---------------------------------------------------------------- Resumen
rs_ = wb.create_sheet('Resumen', 0)
rs_['A1'] = 'GORE Puno · UE 001 Sede Central · mapeo de disponibilidad por proyecto'; rs_['A1'].font = F(bold=True, size=13)
rs_['A2'] = f'Al {FECHA}. PIM por estado situacional (suma de PIM de los proyectos en cada estado). Fórmulas sobre la hoja Base.'; rs_['A2'].font = F(italic=True)
heads = ['Gerencia / Oficina', 'Proyectos (CUI-meta)', 'PIM 2026', 'Certificación', 'Saldo sin certificar', 'Devengado'] + [f'PIM · {s}' for s in STATES]
for j, h in enumerate(heads, 1):
    c = rs_.cell(4, j, h); c.font = F(bold=True, color='FFFFFF'); c.fill = NAVY; c.border = BOX; c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
rs_.row_dimensions[4].height = 36
gers = list(collections.OrderedDict.fromkeys(r['path'][0] for r in rows))
rng = lambda col: f'Base!${col}${BR0+1}:${col}${BLAST}'
for i, g in enumerate(gers):
    n = 5 + i; rs_.cell(n, 1, g)
    rs_.cell(n, 2, f'=COUNTIFS({rng("B")},$A{n})')
    for j, col in zip((3, 4, 6), 'HIK'): rs_.cell(n, j, f'=SUMIFS({rng(col)},{rng("B")},$A{n})')
    rs_.cell(n, 5, f'=C{n}-D{n}')
    for k, s in enumerate(STATES): rs_.cell(n, 7 + k, f'=SUMIFS({rng("H")},{rng("B")},$A{n},{rng("P")},"{s}")')
tn = 5 + len(gers); rs_.cell(tn, 1, 'TOTAL UE 001')
for j in range(2, 12): rs_.cell(tn, j, f'=SUM({L(j)}5:{L(j)}{tn-1})')
for n in range(5, tn + 1):
    for j in range(1, 12):
        c = rs_.cell(n, j); c.font = F(bold=(n == tn)); c.border = BOX
        if j >= 3: c.number_format = NUM
        if n == tn: c.fill = GRY
rs_.column_dimensions['A'].width = 52
for j in range(2, 12): rs_.column_dimensions[L(j)].width = 17
rs_.cell(tn + 2, 1, 'Definición provisional de Estado situacional (sujeta a mejora):').font = F(bold=True)
defs = ['Sin certificar: PIM > 0 y sin certificación → disponible.', 'Solo certificado: con certificación y sin compromiso → posible.',
        'Comprometido: con compromiso anual → no disponible.', 'Devengado total: devengado ≥ PIM.', 'Sin PIM: PIM = 0 (categoría añadida para casos de borde).']
for i, d in enumerate(defs): rs_.cell(tn + 3 + i, 1, d).font = F()
rs_.freeze_panes = 'B5'

# ---------------------------------------------------------------- Control
ct = wb.create_sheet('Control')
ct['A1'] = 'Control y conciliación'; ct['A1'].font = F(bold=True, size=13)
hh = ['Control', 'Calculado (hoja Base)', 'Referencia', 'Diferencia', 'Resultado']
for j, h in enumerate(hh, 1):
    c = ct.cell(3, j, h); c.font = F(bold=True, color='FFFFFF'); c.fill = NAVY; c.border = BOX
ctr = 4
def add(label, formula, ref, tol=0.5):
    global ctr
    ct.cell(ctr, 1, label); ct.cell(ctr, 2, formula); ct.cell(ctr, 3, ref)
    ct.cell(ctr, 4, f'=B{ctr}-C{ctr}'); ct.cell(ctr, 5, f'=IF(ABS(D{ctr})<={tol},"Conciliado","REVISAR")')
    for j in range(1, 6):
        c = ct.cell(ctr, j); c.font = F(); c.border = BOX
        if j in (2, 3, 4): c.number_format = '#,##0.00'
    ctr += 1
add('PIM: Base vs TOTAL UE 001 del XLS SIAF', f'=SUM({rng("H")})', ue_total['pim'])
add('Certificación: Base vs XLS', f'=SUM({rng("I")})', ue_total['cert'])
add('Compromiso anual: Base vs XLS', f'=SUM({rng("J")})', ue_total['comp'])
add('Devengado: Base vs XLS', f'=SUM({rng("K")})', ue_total['dev'])
for nm, ((ws_, tr), nrows) in sheets.items():
    add(f'PIM hoja "{nm}" vs Base', f"='{nm}'!D{tr}", f'=SUMIFS({rng("H")},{rng("A")},"{nm}")')
    add(f'Devengado hoja "{nm}" vs Base', f"='{nm}'!G{tr}", f'=SUMIFS({rng("K")},{rng("A")},"{nm}")')
add('Filas CUI-meta en Base vs metas en XLS', f'=COUNTA({rng("E")})', len(metas), 0)
ctr += 1
ct.cell(ctr, 1, 'Datos faltantes o por validar (informativo)').font = F(bold=True); ctr += 1
nocart = sum(1 for r in rows if not r['car']); sinclas = [r for r in rows if r['p'] is None]
infos = [('Filas sin gerencia asignada (CUI no está en Alex 16/09)', f'=COUNTIFS({rng("A")},"Sin clasificar")'),
         ('Filas sin datos en Cartera PMI (montos 2026-2029 = "No registrado en cartera")', f'=COUNTIFS({rng("Q")},"{NOCART}")'),
         ('Filas con nota de clasificación (meta sin número, repetido, etc.)', f'=COUNTIFS({rng("V")},"?*")'),
         ('Filas con Base legal (sin fuente: en blanco)', f'=COUNTIFS({rng("U")},"?*")')]
for lab, f_ in infos:
    ct.cell(ctr, 1, lab).font = F(); ct.cell(ctr, 2, f_).font = F(); ctr += 1
ctr += 1
ct.cell(ctr, 1, 'CUI sin clasificar:').font = F(bold=True); ctr += 1
for r in sinclas:
    ct.cell(ctr, 1, f"{r['cui']} meta {r['meta']} · PIM {r['pim']:,.0f} · {r['name'][:90]}").font = F(); ctr += 1
ct.column_dimensions['A'].width = 90
for col in 'BCDE': ct.column_dimensions[col].width = 22

# ---------------------------------------------------------------- Modalidad (regla)
md = wb.create_sheet('Regla modalidad')
md['A1'] = 'Regla de Modalidad de ejecución (derivada del SIAF, específica de gasto con PIM > 0, ordenada por monto, separada por " / ")'; md['A1'].font = F(bold=True)
rules = [('"COSTO DE CONSTRUCCION POR CONTRATA" (2.6.2x, 2.6.23)', 'Contrata'), ('"COSTO DE CONSTRUCCION POR ADMINISTRACION DIRECTA - ..."', 'Administración directa'),
         ('"ELABORACION DE EXPEDIENTES TECNICOS" (2.6.81.3.1)', 'Expediente técnico'), ('"ESTUDIO DE PREINVERSION" (2.6.81.2.1)', 'Estudio de preinversión'),
         ('"GASTO POR LAUDOS ARBITRALES..." (2.6.81.4.4)', 'Laudos arbitrales'), ('2.6.3.x (equipos, mobiliario, vehículos)', 'Equipamiento'),
         ('2.6.81.4.1-3 y 2.6.71.6.1-3 (personal, bienes, servicios)', 'Personal/bienes/servicios')]
for i, (a, b) in enumerate(rules):
    md.cell(3 + i, 1, a).font = F(); md.cell(3 + i, 2, b).font = F()
md.cell(3 + len(rules) + 1, 1, 'Si no hay PIM en ninguna específica se usan las que tengan PIA, certificación, compromiso o devengado. "Personal/bienes/servicios" se rotula literal; confirmar con Alex si equivale a administración directa.').font = F(italic=True)
md.column_dimensions['A'].width = 80; md.column_dimensions['B'].width = 28

# orden de hojas
order_ = ['Resumen', 'Gerencia de Infraestructura', 'Gerencias y Oficinas', 'Sin clasificar', 'Base', 'Control', 'Regla modalidad']
wb._sheets = [wb[n] for n in order_ if n in wb.sheetnames]
OUT.parent.mkdir(exist_ok=True)
wb.save(OUT)
print('OK', OUT, len(rows), 'filas;', 'sin clasificar', len(sinclas), '; sin cartera', nocart)
print('PIM', sum_('pim'), 'cert', sum_('cert'), 'comp', sum_('comp'), 'dev', sum_('dev'))
