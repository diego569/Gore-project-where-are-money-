"""Cotejo: (M) mi reporte SINVERSION | (A) Alex-Power Query | (P) practicante Consulta Amigable."""
import re, collections, warnings, pathlib, openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter as L
warnings.filterwarnings('ignore')
ROOT = pathlib.Path(r'E:/New folder/Economia/8vo/gore/Cash data request')
MINE = ROOT / 'outputs/SINVERSION_UE001_2026-10-06.xlsx'
PQ = ROOT / 'outputs/alex_central/Alex_Sede_Central_2026.xlsx'
PRAC = ROOT / 'ref docs/Proyectos Gore-Puno.xlsx'
OUT = ROOT / 'outputs/Cotejo_3_fuentes_2026-10-06.xlsx'
TOL = 1.0   # redondeo de Consulta Amigable (enteros)
isn = lambda v: isinstance(v, (int, float)) and not isinstance(v, bool)
n0 = lambda v: float(v) if isn(v) else 0.0
def meta4(v):
    try: return '%04d' % int(v)
    except Exception: return None
norm = lambda s: re.sub(r'\s+', ' ', str(s or '')).strip().upper()

# ------------------------------------------------ M: mi Base
wb = openpyxl.load_workbook(MINE, data_only=True); ws = wb['Base']
M = []
for r in ws.iter_rows(min_row=5, values_only=True):
    if not r[4]: continue
    NOC = 'No registrado en cartera'
    M.append(dict(sheet=r[0], ger=r[1], fun=r[3], cui=str(r[4]), meta=meta4(r[5]), name=r[6], pim=r[7], cert=r[8], comp=r[9], dev=r[10],
                  mod=r[14], estado=r[15], pmi=[None if (v == NOC or v is None) else v for v in r[16:20]], nota=r[21], grupo=r[2]))
Mk = {(m['cui'], m['meta']): m for m in M}
assert len(Mk) == len(M)

# ------------------------------------------------ A: Alex Power Query
wa = openpyxl.load_workbook(PQ, data_only=True)
CUI = re.compile(r'^\s*(\d{6,7})\s*[:.\s]')
FUN = re.compile(r'^\d{2}\.\s')
A = []
for sn in ['Gerencia de Infraestructura', 'Gerencias y Oficinas']:
    ger = 'GERENCIA DE INFRAESTRUCTURA' if sn.startswith('Gerencia de') else None; fun = None; started = False
    for r in wa[sn].iter_rows(min_row=4, max_col=15, values_only=True):
        a, b, c = r[0], r[1], r[2]
        if isinstance(a, str) and not CUI.match(str(c or '')):
            s = a.strip()
            if s.upper().startswith('TOTAL 001') or s.upper().startswith('TOTAL G'): fun = None; continue
            if s.upper().startswith('METAS SIAF POR CLASIFICAR'): ger = 'SIN CLASIFICAR'; fun = None; continue
            if s.upper().startswith('TOTAL METAS'): continue
            if FUN.match(s): fun = s; continue
            if a.isdigit(): continue
            if sn.startswith('Gerencias') and not s.startswith('001.'): ger = s; fun = None
            continue
        if isinstance(c, str) and CUI.match(c) and isn(r[3]):
            A.append(dict(sheet=sn, ger=ger, fun=fun, cui=CUI.match(c).group(1), meta=meta4(b), name=c.strip(), pim=r[3], cert=r[4], dev=r[6],
                          mod=r[9], estado=r[10], pmi=[(v if isn(v) else None) for v in r[11:15]]))
Ak = collections.defaultdict(list)
for x in A: Ak[(x['cui'], x['meta'])].append(x)
siaf = wa['SIAF']; AS = collections.defaultdict(lambda: [0, 0, 0])
for r in siaf.iter_rows(min_row=7, values_only=True):
    if r[1]:
        for i in range(3): AS[str(r[1])][i] += n0(r[3 + i])

# ------------------------------------------------ P: practicante
wp = openpyxl.load_workbook(PRAC, data_only=True)['Proyectos']
P = {}; lists = collections.defaultdict(list); sect = None
for i, r in enumerate(wp.iter_rows(min_row=16, values_only=True), 16):
    t = str(r[0] or '')
    if t.startswith('Proyectos que'): sect = t.strip(); continue
    m = re.match(r'^(\d{7}):\s*(.*)', t)
    if not m: continue
    if sect: lists[sect].append((m.group(1), m.group(2), n0(r[1]))); continue
    P[m.group(1)] = dict(name=m.group(2), pia=n0(r[1]), pim=n0(r[2]), cert=n0(r[3]), comp=n0(r[4]), dev=n0(r[6]), pmi=r[15])

# ------------------------------------------------ agregados por CUI
def agg(rows, keys):
    d = collections.defaultdict(lambda: dict.fromkeys(keys, 0.0))
    for x in rows:
        for k in keys: d[x['cui']][k] += n0(x.get(k))
    return d
Mc = agg(M, ['pim', 'cert', 'comp', 'dev']); Ac = agg(A, ['pim', 'cert', 'dev'])
for k, v in AS.items(): Ac.setdefault(k, dict(pim=0, cert=0, dev=0))

# ------------------------------------------------ estilos
NAVY = PatternFill('solid', fgColor='002060'); RED = PatternFill('solid', fgColor='F8CBAD'); GRN = PatternFill('solid', fgColor='E2EFDA'); YEL = PatternFill('solid', fgColor='FFF2CC')
FN = lambda size=9, **k: Font(name='Arial', size=size, **k)
out = openpyxl.Workbook(); out.remove(out.active)
def sheet(name, heads, rows, widths=None, numcols=(), flagcol=None, title=None):
    ws = out.create_sheet(name); r0 = 1
    if title: ws['A1'] = title; ws['A1'].font = FN(bold=True, size=11); r0 = 3
    for j, h in enumerate(heads, 1):
        c = ws.cell(r0, j, h); c.font = FN(bold=True, color='FFFFFF'); c.fill = NAVY; c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
    for i, row in enumerate(rows, r0 + 1):
        for j, v in enumerate(row, 1):
            c = ws.cell(i, j, v); c.font = FN()
            if j in numcols: c.number_format = '#,##0.00;(#,##0.00);"-"'
        if flagcol:
            val = str(row[flagcol - 1])
            f = RED if val.startswith(('REVISAR', 'DIFIERE', 'NO ')) else (YEL if val.startswith(('Obs', 'Posible', 'Solo')) else GRN)
            ws.cell(i, flagcol).fill = f
    for j, w in enumerate(widths or [], 1): ws.column_dimensions[L(j)].width = w
    ws.freeze_panes = ws.cell(r0 + 1, 1); ws.auto_filter.ref = f'A{r0}:{L(len(heads))}{r0 + len(rows)}'
    ws.row_dimensions[r0].height = 34
    return ws

# ================= 1) Totales
SIAF_REF = dict(pim=905026984.0, cert=843143127.18, comp=606768924.59, dev=433575772.65)
Pc = {c for c in P if c in Mc}
tot = []
def sm(d, k, keys=None): return sum(v[k] for c, v in d.items() if keys is None or c in keys)
for k, lab in [('pim', 'PIM'), ('cert', 'Certificación'), ('comp', 'Compromiso anual'), ('dev', 'Devengado')]:
    mine = sm(Mc, k); ref = SIAF_REF[k]
    a = sum(x[k] for x in A) if k in ('pim', 'cert', 'dev') else None
    pb = sum(v[k] for v in P.values()); pr = sum(P[c][k] for c in Pc)
    tot.append([lab, ref, mine, a, pb, pr, mine - ref, (a - ref) if a is not None else 'n/d (PQ no trae compromiso)', pr - mine])
sheet('Totales', ['Concepto', 'Referencia XLS SIAF (UE 001)', 'Mi reporte', 'Alex Power Query (hojas Anexo)', 'Practicante: todo el pliego (todas las UEs)',
                  'Practicante: solo CUI que están en UE 001', 'Mi reporte − XLS', 'PQ − XLS', 'Practic. (CUI UE001) − Mi reporte'], tot,
      [22, 20, 18, 20, 22, 22, 16, 22, 20], numcols=range(2, 10),
      title='Totales UE 001 por fuente (al 06/10/2026). El practicante trabaja con TODAS las UEs del pliego; su cifra solo es comparable en los CUI de UE 001, y aun así incluye lo que otras UEs ejecutan en el mismo CUI.')

# ================= 2) CUI 3 fuentes
cuis = sorted(set(Mc) | set(Ac) | set(P))
rows = []; stat = collections.Counter()
for c in cuis:
    m = Mc.get(c); a = Ac.get(c); p = P.get(c)
    nm = next((x['name'] for x in M if x['cui'] == c), None) or next((x['name'] for x in A if x['cui'] == c), None) or (p['name'] if p else '')
    flags = []
    mp = m and m['pim'];
    if m and a:
        for k in ('pim', 'cert', 'dev'):
            if abs(m[k] - a[k]) > 0.5: flags.append(f'Mío≠PQ en {k}')
    if m is None: flags.append('Solo en PQ' if a and not p else 'No está en mi reporte')
    if a is None and m: flags.append('No está en PQ')
    if p and m:
        for k in ('pim', 'cert', 'comp', 'dev'):
            d = p[k] - m[k]
            if d < -TOL: flags.append(f'REVISAR: practic. < mío en {k}')
            elif d > TOL: flags.append(f'Posible otras UEs ({k})')
    if m and not p: flags.append('No está en práctic. (PIM>0?)' if m['pim'] > 0 else 'Sin PIM; no figura en práctic.')
    if p and not m: flags.append('Solo práctic. (otra UE)')
    flag = ' | '.join(flags) or 'OK'
    stat['OK' if flag == 'OK' else ('REVISAR' if any(f.startswith('REVISAR') or f.startswith('Mío≠PQ') for f in flags) else 'Obs')] += 1
    g = lambda d, k: d[k] if d and k in d else None
    rows.append([c, nm, g(m, 'pim'), g(a, 'pim'), g(p, 'pim'), g(m, 'cert'), g(a, 'cert'), g(p, 'cert'), g(m, 'comp'), g(p, 'comp'), g(m, 'dev'), g(a, 'dev'), g(p, 'dev'), flag])
sheet('CUI_3_fuentes', ['CUI', 'Proyecto', 'PIM Mío', 'PIM PQ', 'PIM Prac.', 'Cert. Mío', 'Cert. PQ', 'Cert. Prac.', 'Comp. Mío', 'Comp. Prac.', 'Dev. Mío', 'Dev. PQ', 'Dev. Prac.', 'Resultado'],
      rows, [9, 50] + [14] * 11 + [60], numcols=range(3, 14), flagcol=14,
      title='Cotejo por CUI (suma de metas). Tolerancia 1 sol por redondeo del practicante; "Posible otras UEs" = práctico mayor, esperable si otras UEs ejecutan el mismo CUI.')

# ================= 3) Fila CUI-meta: mío vs PQ
rows = []; cnt = collections.Counter()
for k in sorted(set(Mk) | set(Ak)):
    m = Mk.get(k); al = Ak.get(k, []); a = al[0] if al else None; fl = []
    if m and not a: fl.append('NO está en PQ (CUI-meta)')
    if a and not m: fl.append('NO está en mi reporte')
    if len(al) > 1: fl.append('PQ repite CUI-meta')
    if m and a:
        for kk in ('pim', 'cert', 'dev'):
            if abs(n0(m[kk]) - n0(a[kk])) > 0.5: fl.append(f'DIFIERE {kk}')
        mg = norm(m['ger']).replace('SIN CLASIFICAR (PENDIENTE DE ASIGNAR GERENCIA)', 'SIN CLASIFICAR')
        if mg != norm(a['ger']): fl.append('DIFIERE gerencia')
        if norm(m['fun']) != norm(a['fun']) and 'SIN CLASIFICAR' not in mg: fl.append('Obs función distinta')
        for i in range(4):
            x, y = m['pmi'][i], a['pmi'][i]
            if (x is None) != (y is None) or (x is not None and abs(x - y) > 0.5): fl.append(f'DIFIERE PMI {2026 + i}'); break
    for f in fl: cnt[f] += 1
    if not fl: cnt['OK'] += 1
    rows.append([k[0], k[1], (m or a)['name'], m and m['pim'], a and a['pim'], m and m['ger'], a and a['ger'], m and m['fun'], a and a['fun'],
                 m and m['estado'], a and a['estado'], m and m['mod'], a and a['mod'], ' | '.join(fl) or 'OK'])
sheet('Mio_vs_PQ_filas', ['CUI', 'Meta', 'Proyecto', 'PIM Mío', 'PIM PQ', 'Gerencia Mío', 'Gerencia PQ', 'Función Mío', 'Función PQ', 'Estado Mío', 'Estado PQ', 'Modalidad Mío', 'Modalidad PQ', 'Resultado'],
      rows, [9, 7, 46, 14, 14, 36, 36, 24, 24, 16, 14, 30, 30, 50], numcols=(4, 5), flagcol=14, title='Cotejo fila a fila (CUI + meta): mi reporte vs Alex Power Query (hojas Gerencia de Infraestructura y Gerencias y Oficinas).')

# ================= 4) Estado vs listas del practicante
rows = []
def state_cui(c):
    ms = [x for x in M if x['cui'] == c]
    return ' / '.join(sorted({x['estado'] for x in ms})) if ms else 'No está en UE 001', sum(n0(x['pim']) for x in ms), any(x['pmi'][0] is None and x['pmi'][1] is None for x in ms)
expect = {'Proyectos que NO-PMI': None, 'Proyectos que NO-Cert-Comp': 'Sin certificar', 'Proyectos que Solo Certificaron': 'Solo certificado'}
for sect, items in lists.items():
    key = next(k for k in expect if sect.startswith(k[:20]))
    for c, nm, pim in items:
        st, mpim, nopmi = state_cui(c)
        if st == 'No está en UE 001': res = 'Posible otra UE (no verificable)'
        elif expect[key] is None: res = 'OK' if nopmi else 'DIFIERE: en cartera según mi reporte'
        else: res = 'OK' if st == expect[key] else f'DIFIERE: mío = {st}'
        rows.append([sect, c, nm, pim, mpim, st, 'Sí' if nopmi else 'No', res])
sheet('Listas_practicante', ['Lista del practicante', 'CUI', 'Proyecto', 'PIM (práct.)', 'PIM (mío, UE 001)', 'Estado mío', 'Sin cartera (mío)', 'Resultado'], rows,
      [30, 9, 50, 14, 14, 24, 12, 38], numcols=(4, 5), flagcol=8, title='Listas manuales del practicante vs el estado situacional calculado en mi reporte.')
# inversa: mis Sin certificar / Solo certificado con PIM>0 que no están en listas
inlist = {c for items in lists.values() for c, _, _ in items}
rev = collections.defaultdict(lambda: [0, 0.0])
for m in M:
    if m['estado'] in ('Sin certificar', 'Solo certificado') and n0(m['pim']) > 0 and m['cui'] not in inlist:
        rev[m['estado']][0] += 1; rev[m['estado']][1] += n0(m['pim'])

# ================= 5) Resumen y hallazgos
ws = out.create_sheet('Resumen', 0)
ws['A1'] = 'Cotejo de tres fuentes · UE 001 Sede Central · 06/10/2026'; ws['A1'].font = FN(bold=True, size=13)
lines = [
    ('Fuentes', ''),
    ('Mío', 'SINVERSION_UE001_2026-10-06.xlsx (XLS SIAF + Cartera + agrupación Alex 16/09), %d filas CUI-meta, %d CUI' % (len(M), len(Mc))),
    ('Alex PQ', 'Alex_Sede_Central_2026.xlsx (Power Query), %d filas CUI-meta en anexos, %d CUI' % (len(A), len({x['cui'] for x in A}))),
    ('Practicante', 'Proyectos Gore-Puno.xlsx (Consulta Amigable, TODO el pliego 458), %d CUI' % len(P)),
    ('', ''),
    ('Resultados', ''),
    ('Total UE 001: mío vs PQ', 'PIM, certificación y devengado: ver hoja Totales (ambos coinciden con el XLS SIAF).'),
    ('Filas CUI-meta', 'Iguales en las 3 magnitudes y misma ubicación: %d de %d (detalle en Mio_vs_PQ_filas)' % (cnt['OK'], len(set(Mk) | set(Ak)))),
]
for k, v in sorted(cnt.items(), key=lambda kv: -kv[1]):
    if k != 'OK': lines.append(('  · ' + k, f'{v} filas'))
lines += [('', ''), ('CUI (hoja CUI_3_fuentes)', 'OK: %d · con observación: %d · a revisar: %d' % (stat['OK'], stat['Obs'], stat['REVISAR'])), ('', ''),
          ('Listas del practicante', 'ver hoja Listas_practicante (%d CUI)' % sum(len(v) for v in lists.values()))]
for s, (n, p_) in rev.items(): lines.append((f'  · Mío "{s}" con PIM>0 fuera de sus listas', f'{n} filas, PIM {p_:,.0f}'))
lines += [('', ''), ('Cobertura de campos', ''),
          ('Estado situacional', f'PQ: "Por completar" en {sum(1 for x in A if x["estado"] == "Por completar")} de {len(A)} filas; mío: calculado en {len(M)}.'),
          ('Modalidad', f'PQ: vacía o texto de grupo Alex; mío: derivada del SIAF.'),
          ('Compromiso anual', 'PQ no lo trae; mío y practicante sí.')]
for i, (k, v) in enumerate(lines, 3):
    ws.cell(i, 1, k).font = FN(bold=(v == '')); ws.cell(i, 2, v).font = FN()

ncomp = [c for c in Mc if c in P and abs(P[c]['comp'] - Mc[c]['comp']) > TOL and all(abs(P[c][k] - Mc[c][k]) <= TOL for k in ('pim', 'cert', 'dev'))]
nopr = [c for c in Mc if c not in P]
lines += [('', ''), ('Hallazgos', ''),
  ('1. Mío vs Power Query', 'Coinciden en PIM, certificación, devengado, ubicación y montos PMI en %d de %d filas. Única diferencia: CUI 2194989 (PIM 133,512): yo lo ubiqué en Gerencia de Infraestructura (Alex lo tiene sin número de meta) y PQ lo deja en "Metas por clasificar".' % (cnt['OK'], len(set(Mk) | set(Ak)))),
  ('2. Practicante incompleto', 'El archivo dice "Resultados 1-400 de 447, Página 1 de 2": solo se descargó la página 1. Faltan 47 proyectos; %d CUI de UE 001 (PIM %s) no aparecen en su base.' % (len(nopr), format(sum(Mc[c]['pim'] for c in nopr), ',.0f'))),
  ('3. Practicante mezcla UEs', 'Su base es de todo el pliego 458 (PIM 971.6 M cargado, 987.99 M según su cabecera); no permite aislar UE 001. %d CUI son de otras UEs.' % sum(1 for c in P if c not in Mc)),
  ('4. Compromiso anual', 'En %d CUI coinciden PIM/certificación/devengado pero el compromiso anual difiere entre el XLS SIAF y la Consulta Amigable, en ambos sentidos. Caso crítico: CUI 2357091 (PIM 10.4 M): XLS compromiso 0 (mío = "Solo certificado"), Consulta Amigable compromiso 9.6 M (= "Comprometido"). Causa no determinada (¿hora de extracción o definición?). Afecta el estado.' % len(ncomp)),
  ('5. Estado por meta vs por CUI', 'Mi estado es por CUI-meta; el practicante clasifica por CUI. CUI con varias metas pueden salir con estados distintos (ej. 2301117, 2001621).'),
  ('6. Listas del practicante', 'NO-PMI: 18 de 20 coinciden con "sin cartera" en mi reporte; NO-Cert-Comp: 17 de 19; Solo Certificaron: 18 de 18. Las no coincidentes son CUI de otras UEs (2313145, 2714101, 2325694, 2393101).'),
  ('7. PQ no completa campos', 'Estado situacional = "Por completar" en todas las filas; Modalidad vacía o con texto del grupo de Alex; sin compromiso anual.')]
for i, (k, v) in enumerate(lines[len(lines) - 9:], 3 + len(lines) - 9):
    ws.cell(i, 1, k).font = FN(bold=(v == '')); ws.cell(i, 2, v).font = FN(); ws.cell(i, 2).alignment = Alignment(wrap_text=True, vertical='top')

ws.column_dimensions['A'].width = 48; ws.column_dimensions['B'].width = 120
out.save(OUT)
print('saved', OUT)
print('cnt', cnt); print('stat', stat); print('rev', dict(rev))
print(len(M), len(A), len(P), {k: len(v) for k, v in lists.items()})
