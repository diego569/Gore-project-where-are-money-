"""Compara el Excel generado (hoja Base) con el Excel del practicante ("Proyectos Gore-Puno Completo-OFICIAL.xlsx").

Uso:  python validar_listas.py "<practicante.xlsx>" [<SINVERSION_PROYECTOS_...xlsx>] [<salida.xlsx>]
Genera un Excel con: Resumen (por lista del practicante), Detalle (CUI por CUI), Cifras (PIM/certificación/compromiso/devengado) y
"Solo en nuestro Excel" (proyectos que mostramos como con saldo y el practicante no listó).
"""
import collections, pathlib, re, sys
import openpyxl, xlsxwriter

RAIZ = pathlib.Path(__file__).resolve().parent
LISTAS = {'A': 'NO está en el PMI (sin movimiento)', 'G': 'NO está en el PMI (con movimiento)', 'O': 'Sí está en el PMI · sin certificar ni comprometer',
          'R': 'Sí está en el PMI · solo certificaron'}


def leer_practicante(p):
    wb = openpyxl.load_workbook(p, data_only=True)
    ws = wb['RESUMEN']
    listas = collections.defaultdict(list)
    for row in ws.iter_rows():
        for c in row:
            v = c.value
            if isinstance(v, str):
                m = re.match(r'^(\d{7}):', v.strip())
                if m and c.column_letter in LISTAS:
                    nxt = ws.cell(c.row, c.column + 1).value
                    listas[c.column_letter].append((m.group(1), v.strip(), nxt if isinstance(nxt, (int, float)) else None))
    cifras = {}
    for nm in wb.sheetnames:
        if nm.startswith('Ordenado'):
            for row in wb[nm].iter_rows(values_only=True):
                if row and isinstance(row[0], str):
                    m = re.match(r'^(\d{7}):', row[0].strip())
                    if m:
                        n = [x if isinstance(x, (int, float)) else None for x in row[1:9]]
                        cifras[m.group(1)] = dict(pia=n[0], pim=n[1], cert=n[2], comp=n[3], dev=n[5])
    return listas, cifras


def leer_base(p):
    ws = openpyxl.load_workbook(p, data_only=True)['Base']
    rows = list(ws.iter_rows(min_row=4, values_only=True)); h = rows[0]
    por = collections.defaultdict(list)
    for r in rows[1:]:
        if r[0]:
            d = dict(zip(h, r)); por[str(d['CUI'])].append(d)
    return por


def num(x):
    return x if isinstance(x, (int, float)) else 0.0


def main():
    prac = sys.argv[1]
    base_p = sys.argv[2] if len(sys.argv) > 2 else None
    if not base_p:
        cands = sorted((RAIZ / 'salida').glob('*/SINVERSION_PROYECTOS_*.xlsx'), key=lambda p: p.stat().st_mtime)
        base_p = str(cands[-1])
    out = sys.argv[3] if len(sys.argv) > 3 else str(RAIZ.parent / 'outputs' / ('Validacion_vs_practicante_' + pathlib.Path(base_p).stem[-10:] + '.xlsx'))
    listas, cifras = leer_practicante(prac); mine = leer_base(base_p)
    wb = xlsxwriter.Workbook(out)
    F = lambda **k: wb.add_format({'font_name': 'Arial', 'font_size': 9, **k})
    hd = F(bold=True, font_color='#FFFFFF', bg_color='#002060', text_wrap=True, align='center', valign='vcenter', border=1)
    tx = F(); wr = F(text_wrap=True, valign='top'); nm = F(num_format='#,##0;(#,##0);"-"'); ti = F(bold=True, font_size=12); nt = F(italic=True, font_color='#595959')
    sems = ['Disponible', 'Con condiciones', 'Rebajando certificación', 'Restringido', 'No mover', 'Sin saldo útil', 'Sin PIM', 'Bolsa genérica', 'No está en nuestra Base']

    def mio(cui):
        rs = mine.get(cui)
        if not rs:
            return None
        r = max(rs, key=lambda x: num(x['PIM 2026']))
        return dict(sem=r['Semáforo'], disp=sum(num(x['Disponible legal (S/)']) for x in rs), pot=sum(num(x['Potencial rebajando certificación (S/)']) for x in rs),
                    pim=sum(num(x['PIM 2026']) for x in rs), cert=sum(num(x['Certificación']) for x in rs), motivo=r['Motivo / condiciones'], tipo=r['Tipo de inversión'],
                    pmi26=r['PMI 2026 (S/)'], av=r['% Avance financiero acumulado'], prog=r['Programación 12-B 2026 (total)'], ue=r['UE'])

    # ---- Resumen
    ws = wb.add_worksheet('Resumen')
    ws.write(0, 0, 'Validación: Excel automático vs. listas del practicante', ti)
    ws.write(1, 0, 'Cada fila es una lista del practicante; las columnas muestran en qué semáforo cayó cada uno de esos proyectos en nuestro Excel.', nt)
    ws.write_row(3, 0, ['Lista del practicante', 'Proyectos', 'PIM (practicante)'] + sems, hd); ws.set_row(3, 40)
    for i, (k, nombre) in enumerate(LISTAS.items()):
        c = collections.Counter(); pim = 0.0
        for cui, _t, v in listas.get(k, []):
            m = mio(cui); c[m['sem'] if m else 'No está en nuestra Base'] += 1; pim += num(v)
        ws.write(4 + i, 0, f'{k}. {nombre}', tx); ws.write(4 + i, 1, len(listas.get(k, [])), tx); ws.write(4 + i, 2, pim, nm)
        for j, s in enumerate(sems):
            ws.write(4 + i, 3 + j, c[s], tx)
    ws.set_column(0, 0, 58); ws.set_column(1, 2, 14); ws.set_column(3, 12, 14)
    ws.write(10, 0, 'Cómo leerlo', F(bold=True))
    for i, t in enumerate([
            'A y G (no están en el PMI): esperamos "Con condiciones" o "Disponible" (salvo PROCOMPITE = Restringido, o proyectos con avance > 90% = No mover).',
            'O (en el PMI, sin certificar ni comprometer): el practicante los considera plata; nosotros no, porque el 12-B y el PMI 2026 los respaldan ("Sin saldo útil").',
            'R (en el PMI, solo certificaron): "Rebajando certificación" cuando lo certificado supera lo programado en el 12-B; el resto queda "Sin saldo útil".',
            'Los motivos de cada caso están en la hoja Detalle.']):
        ws.write(11 + i, 0, t, tx)

    # ---- Detalle
    wd = wb.add_worksheet('Detalle')
    H = ['Lista', 'CUI', 'Proyecto (practicante)', 'PIM practicante', 'Nuestro semáforo', 'Disponible legal', 'Potencial rebajando cert.', 'PIM (nuestro)', 'Tipo', 'PMI 2026', '12-B 2026', '% avance', 'Motivo (nuestro)']
    wd.write_row(0, 0, H, hd); wd.set_row(0, 30); r = 1
    for k in LISTAS:
        for cui, texto, v in listas.get(k, []):
            m = mio(cui)
            fila = [k, cui, texto[9:], num(v)] + ([m['sem'], m['disp'], m['pot'], m['pim'], m['tipo'], m['pmi26'], m['prog'], m['av'], m['motivo']] if m else ['No está en nuestra Base'] + [''] * 8)
            for j, x in enumerate(fila):
                if isinstance(x, (int, float)): wd.write_number(r, j, x, F(num_format='0%') if j == 11 else nm)
                else: wd.write_string(r, j, str(x), wr if j in (2, 12) else tx)
            r += 1
    wd.autofilter(0, 0, r - 1, len(H) - 1); wd.freeze_panes(1, 3)
    for j, w in enumerate([6, 10, 60, 14, 22, 14, 16, 14, 18, 14, 14, 9, 70]):
        wd.set_column(j, j, w)

    # ---- Cifras
    wc = wb.add_worksheet('Cifras')
    wc.write_row(0, 0, ['CUI', 'PIM practicante', 'PIM nuestro', 'Dif. PIM', 'Cert. practicante', 'Cert. nuestra', 'Dif. cert.', 'Comp. practicante', 'Comp. nuestro', 'Dif. comp.', 'Dev. practicante', 'Dev. nuestro', 'Dif. dev.'], hd)
    r = 1; difs = 0
    for cui, p in sorted(cifras.items()):
        rs = mine.get(cui)
        if not rs: continue
        n = dict(pim=sum(num(x['PIM 2026']) for x in rs), cert=sum(num(x['Certificación']) for x in rs), comp=sum(num(x['Compromiso anual']) for x in rs), dev=sum(num(x['Devengado']) for x in rs))
        vals = [cui]
        hay = False
        for k in ('pim', 'cert', 'comp', 'dev'):
            a, b = num(p[k]), n[k]; vals += [a, b, a - b]; hay |= abs(a - b) > 1
        for j, x in enumerate(vals):
            if isinstance(x, str): wc.write_string(r, j, x, tx)
            else: wc.write_number(r, j, x, nm)
        difs += hay; r += 1
    wc.autofilter(0, 0, r - 1, 12); wc.freeze_panes(1, 1); wc.set_column(0, 0, 10); wc.set_column(1, 12, 15)
    wc.write(r + 1, 0, f'CUI con alguna diferencia mayor a S/ 1: {difs} de {r - 1} comparados', F(bold=True))
    en_prac = {c for l in listas.values() for c, _t, _v in l}

    # ---- Solo en nuestro Excel
    wo = wb.add_worksheet('Solo en nuestro Excel')
    wo.write_row(0, 0, ['CUI', 'UE', 'Proyecto', 'Semáforo', 'Disponible legal', 'Potencial rebajando cert.', 'Tipo', 'Motivo'], hd)
    r = 1
    for cui, rs in sorted(mine.items()):
        if cui in en_prac: continue
        for x in rs:
            if x['Semáforo'] in ('Disponible', 'Con condiciones', 'Rebajando certificación'):
                for j, v in enumerate([cui, x['UE'], x['Proyecto'], x['Semáforo'], num(x['Disponible legal (S/)']), num(x['Potencial rebajando certificación (S/)']), x['Tipo de inversión'], x['Motivo / condiciones']]):
                    if isinstance(v, (int, float)): wo.write_number(r, j, v, nm)
                    else: wo.write_string(r, j, str(v), wr if j in (2, 7) else tx)
                r += 1
    wo.autofilter(0, 0, max(r - 1, 1), 7); wo.freeze_panes(1, 3)
    for j, w in enumerate([10, 6, 70, 22, 15, 16, 20, 80]):
        wo.set_column(j, j, w)
    wb.close()
    print('Listo:', out)


if __name__ == '__main__':
    main()
