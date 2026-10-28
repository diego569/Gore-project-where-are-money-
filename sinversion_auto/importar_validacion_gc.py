"""Carga al maestro (hoja ClasifGC) la clasificación de gasto corriente validada por el GORE.
Uso: python importar_validacion_gc.py <Propuesta_validada.xlsx>
Regla por fila (hoja "Metas"): Validación = Conforme -> se usa "Unidad propuesta"; = Corregir -> se usa "Unidad correcta".
Las filas "No sé", vacías o con unidad "POR DEFINIR" no se cargan (quedan pendientes).
"""
import sys, pathlib, openpyxl
from sinv.maestro import asegurar_estructura
from openpyxl.styles import PatternFill

RAIZ = pathlib.Path(__file__).resolve().parent
src = pathlib.Path(sys.argv[1])
mpath = RAIZ / 'maestro.xlsx'
asegurar_estructura(mpath)
ws = openpyxl.load_workbook(src, data_only=True)['Metas']
cab = [c.value for c in ws[1]]; ix = {n: i for i, n in enumerate(cab)}
wb = openpyxl.load_workbook(mpath); gc = wb['ClasifGC']
fila = {(str(r[0].value or '001').zfill(3), str(r[1].value).zfill(4)): r[0].row for r in gc.iter_rows(min_row=2) if r[1].value}
nuevas = act = omit = 0
for r in ws.iter_rows(min_row=2, values_only=True):
    meta = str(r[ix['Meta']]).zfill(4); val = (r[ix['Validación']] or '').strip()
    unidad = (r[ix['Unidad propuesta']] if val == 'Conforme' else r[ix['Unidad correcta']] if val == 'Corregir' else None)
    unidad = (unidad or '').strip()
    if not unidad or unidad.upper().startswith('POR DEFINIR'):
        omit += 1; continue
    k = ('001', meta)
    if k in fila:
        c = gc.cell(fila[k], 3); c.value = unidad; gc.cell(fila[k], 6).value = 'Validado GORE'; c.fill = PatternFill(fill_type=None); act += 1
    else:
        gc.append(['001', meta, unidad, str(r[ix['Función']] or ''), str(r[ix['Actividad (SIAF)']] or '')[:200], 'Validado GORE']); nuevas += 1
wb.save(mpath)
print(f'ClasifGC: {act} metas actualizadas, {nuevas} agregadas, {omit} sin cargar (sin validar o "POR DEFINIR").')
