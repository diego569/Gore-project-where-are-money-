"""Lector de la Cartera PMI (Banco de Inversiones, .xlsx)."""
import re, datetime
import openpyxl

NOMBRES = {'cui': 'Código Único', 'nombre': 'Nombre de inversión', 'tipo': 'Tipo inversión', 'funcion': 'Función',
           'costo': 'Costo actualizado (S/)', 'dev_acum': 'Devengado acumulado (S/)', 'pim': 'PIM 2026 (S/)',
           'm2026': 'Monto Año 2026 (S/)', 'm2027': 'Monto Año 2027 (S/)', 'm2028': 'Monto Año 2028 (S/)', 'm2029': 'Monto Año 2029 (S/)'}


def _num(v):
    if v in (None, ''):
        return None
    try:
        return float(str(v).replace(',', ''))
    except ValueError:
        return None


def es_cartera(path):
    try:
        wb = openpyxl.load_workbook(path, read_only=True)
        try:
            ws = wb.active
            return any('Reporte de cartera' in str(c) for row in ws.iter_rows(max_row=3, values_only=True) for c in row if c)
        finally:
            wb.close()
    except Exception:
        return False


def leer_cartera(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = list(wb.active.iter_rows(values_only=True))
    finally:
        wb.close()
    fecha = None
    for r in rows[:6]:
        for c in r:
            m = re.search(r'Fecha:\s*(\d{2})/(\d{2})/(\d{4})\s+(\d{1,2}):(\d{2}):(\d{2})\s*([ap])', str(c or ''))
            if m:
                d, mo, y, hh, mi, ss, ap = m.groups(); hh = int(hh) % 12 + (12 if ap == 'p' else 0)
                fecha = datetime.datetime(int(y), int(mo), int(d), hh, int(mi), int(ss))
    hi = next(i for i, r in enumerate(rows) if r and any(str(c).strip() == 'Código Único' for c in r if c))
    head = [str(c or '').strip() for c in rows[hi]]
    idx = {}
    for k, nom in NOMBRES.items():
        j = next((j for j, h in enumerate(head) if h.startswith(nom)), None)
        if j is None:
            raise ValueError(f'La Cartera no tiene la columna "{nom}". ¿Cambió el formato del reporte?')
        idx[k] = j
    out = {}
    for r in rows[hi + 1:]:
        cui = str(r[idx['cui']] or '').strip()
        if not cui:
            continue
        d = {k: r[j] for k, j in idx.items()}
        for k in ('costo', 'dev_acum', 'pim', 'm2026', 'm2027', 'm2028', 'm2029'):
            d[k] = _num(d[k])
        d['cui'] = cui
        out[cui] = d
    return dict(cartera=out, fecha=fecha)
