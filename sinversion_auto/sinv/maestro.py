"""maestro.xlsx: lo que NO viene de los reportes y debe conservarse día a día.

Hojas:
  Clasificacion : en qué gerencia / grupo / función va cada proyecto de la UE 001 (parte de Alex 16/09).
                  Los proyectos nuevos se agregan solos con Gerencia vacía: basta completarla.
  Complementos  : Base legal y Observación por UE+CUI+meta (se escriben a mano).
  Funciones     : código de función -> nombre.
"""
import re, pathlib, warnings, collections
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

warnings.filterwarnings('ignore')
CUIRE = re.compile(r'^(\d{6,7})\s*[:.]\s*(.*)')
FUNRE = re.compile(r'^(\d{2})\s*[.:]\s*(.+)')
HEAD_CLAS = ['UE', 'CUI', 'Meta', 'Hoja', 'Gerencia', 'Grupo', 'Función', 'Orden', 'Origen', 'Proyecto (referencia)']
HEAD_COMP = ['UE', 'CUI', 'Meta', 'Base legal', 'Observación']
INFRA = 'GERENCIA DE INFRAESTRUCTURA'
HOJA_INFRA = 'Gerencia de Infraestructura'
HOJA_GER = 'Gerencias y Oficinas'


def _meta4(v):
    try:
        return '%04d' % int(v)
    except Exception:
        return ''


def sembrar_desde_alex(xlsm_path, xlsm_extra=None):
    """Lee el SINVERSION de Alex (UE 001) y devuelve (clasificación, funciones)."""
    ax = openpyxl.load_workbook(xlsm_path, data_only=True)
    clas = []; funcs = {}; orden = 0

    def reg_fun(s):
        m = FUNRE.match(s.strip())
        if m:
            funcs[m.group(1)] = re.sub(r'\s+', ' ', m.group(2)).strip()
            return f"{m.group(1)}. {funcs[m.group(1)]}"
        return None

    ws = ax[HOJA_INFRA]; grp = None; fun = None
    for a, b, c in ws.iter_rows(min_row=5, max_row=904, max_col=3, values_only=True):
        if isinstance(c, str) and CUIRE.match(c.strip()):
            mm = CUIRE.match(c.strip()); orden += 1
            if grp and fun:
                clas.append(['001', mm.group(1), _meta4(b), HOJA_INFRA, INFRA, ' > '.join(grp), fun, orden, 'Alex 16/09', mm.group(2)])
            continue
        if isinstance(a, str):
            s = a.strip()
            f = reg_fun(s) if re.match(r'^\d{2}\.\s', s) else None
            if f:
                fun = f
            elif s.startswith('001.') or s.upper().startswith(('GERENCIA DE INFRA', 'TOTAL')) or s == 'EJECUCIÓN POR ADMINISTRACIÓN DIRECTA':
                pass
            else:
                grp = (('EJECUCIÓN POR ADMINISTRACIÓN DIRECTA', s) if s.startswith('EJECUCIÓN POR ADMINISTRACIÓN DIRECTA - ') else (s,)); fun = None

    ws = ax[HOJA_GER]; ger = None; fun = None
    for a, b, c in ws.iter_rows(min_row=6, max_row=313, max_col=3, values_only=True):
        if isinstance(c, str) and CUIRE.match(c.strip()):
            mm = CUIRE.match(c.strip()); orden += 1
            clas.append(['001', mm.group(1), _meta4(b), HOJA_GER, ger, '', fun or '', orden, 'Alex 16/09', mm.group(2)])
            continue
        if isinstance(a, str):
            s = a.strip()
            if re.match(r'^\d{2}\.\s', s):
                fun = reg_fun(s)
            elif not s.upper().startswith('TOTAL'):
                ger = s; fun = None
    # nombres de función de todas las hojas de Alex (también de otras UE)
    for w in ax.worksheets:
        if w.title in ('Hoja9', 'Hoja4'):
            continue
        for r in w.iter_rows(max_col=2, values_only=True):
            for v in r:
                if isinstance(v, str) and re.match(r'^\d{2}\s*[.:]\s*[A-ZÁÉÍÓÚÑ]', v.strip()):
                    m = FUNRE.match(v.strip()); funcs.setdefault(m.group(1), re.sub(r'\s+', ' ', m.group(2)).strip())
    if xlsm_extra:
        try:
            ws = openpyxl.load_workbook(xlsm_extra, data_only=True, read_only=True)['por funcion saldo al expediente']
            for r in ws.iter_rows(max_col=3, values_only=True):
                for v in r:
                    if isinstance(v, str) and re.match(r'^\d{2}\s*[.:]\s*[A-ZÁÉÍÓÚÑ]', v.strip()):
                        m = FUNRE.match(v.strip()); funcs.setdefault(m.group(1), re.sub(r'\s+', ' ', m.group(2)).strip())
        except Exception:
            pass
    return clas, funcs


def _estilo_cabecera(ws, heads):
    ws.append(heads)
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='002060'); c.alignment = Alignment(wrap_text=True)
    ws.freeze_panes = 'A2'


def guardar_maestro(path, clas, comp, funcs):
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = 'Clasificacion'
    _estilo_cabecera(ws, HEAD_CLAS)
    for r in clas:
        ws.append(r)
    for col, w in zip('ABCDEFGHIJ', [6, 10, 7, 28, 44, 52, 34, 8, 14, 60]):
        ws.column_dimensions[col].width = w
    ws.auto_filter.ref = f'A1:J{max(2, ws.max_row)}'
    gers = sorted({r[4] for r in clas if r[4]})
    wl = wb.create_sheet('Listas'); wl.append(['Gerencias / oficinas', 'Hojas', 'Grupos (solo Infraestructura)'])
    grupos = sorted({r[5] for r in clas if r[3] == HOJA_INFRA and r[5]})
    for i in range(max(len(gers), len(grupos), 2)):
        wl.append([gers[i] if i < len(gers) else None, [HOJA_INFRA, HOJA_GER][i] if i < 2 else None, grupos[i] if i < len(grupos) else None])
    wl.column_dimensions['A'].width = 50; wl.column_dimensions['B'].width = 30; wl.column_dimensions['C'].width = 60
    dv = DataValidation(type='list', formula1=f"=Listas!$A$2:$A${len(gers) + 1}", allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv); dv.add('E2:E5000')
    dv2 = DataValidation(type='list', formula1=f"=Listas!$C$2:$C${len(grupos) + 1}", allow_blank=True, showErrorMessage=False)
    ws.add_data_validation(dv2); dv2.add('F2:F5000')
    wc = wb.create_sheet('Complementos'); _estilo_cabecera(wc, HEAD_COMP)
    for r in comp:
        wc.append(r)
    for col, w in zip('ABCDE', [6, 10, 7, 40, 60]):
        wc.column_dimensions[col].width = w
    wf = wb.create_sheet('Funciones'); _estilo_cabecera(wf, ['Código', 'Nombre'])
    for k in sorted(funcs):
        wf.append([k, funcs[k]])
    wf.column_dimensions['B'].width = 60
    wb.save(path)


def cargar_maestro(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    clas = []
    for r in wb['Clasificacion'].iter_rows(min_row=2, values_only=True):
        if r[1]:
            r = list(r) + [None] * (10 - len(r))
            clas.append(dict(ue=str(r[0] or '001').zfill(3), cui=str(r[1]).strip(), meta=_meta4(r[2]) if r[2] not in (None, '') else '',
                             hoja=r[3] or '', gerencia=(r[4] or '').strip(), grupo=(r[5] or '').strip(), funcion=(r[6] or '').strip(),
                             orden=r[7] or 10 ** 6, origen=r[8] or '', nombre=r[9] or ''))
    comp = {}
    for r in wb['Complementos'].iter_rows(min_row=2, values_only=True):
        if r[1]:
            comp[(str(r[0] or '001').zfill(3), str(r[1]).strip(), _meta4(r[2]) if r[2] not in (None, '') else '')] = dict(base_legal=r[3] or '', obs=r[4] or '')
    funcs = {str(r[0]).zfill(2): r[1] for r in wb['Funciones'].iter_rows(min_row=2, values_only=True) if r[0]}
    return dict(clas=clas, comp=comp, funcs=funcs)


def agregar_pendientes(path, filas):
    """filas: [[UE, CUI, Meta, Hoja, Gerencia, Grupo, Función, Orden, Origen, Nombre]] con Gerencia vacía."""
    if not filas:
        return 0
    wb = openpyxl.load_workbook(path); ws = wb['Clasificacion']
    existentes = {(str(r[0] or '001').zfill(3), str(r[1]), _meta4(r[2]) if r[2] not in (None, '') else '') for r in ws.iter_rows(min_row=2, values_only=True) if r[1]}
    n = 0
    for f in filas:
        if (f[0], f[1], f[2]) in existentes:
            continue
        ws.append(f); n += 1
        for c in ws[ws.max_row][4:6]:
            c.fill = PatternFill('solid', fgColor='FFF2CC')
    wb.save(path)
    return n


def ubicar(maestro, ue, cui, meta, usados):
    """Devuelve (fila_maestro | None, nota). Misma lógica de coincidencia que la versión 1."""
    ps = [p for p in maestro['clas'] if p['ue'] == ue and p['cui'] == cui]
    if not ps:
        return None, 'Proyecto nuevo: no figura en el maestro'
    libre = lambda p: id(p) not in usados
    for cond, nota in ((lambda p: p['meta'] == meta, ''), (lambda p: p['meta'] == '', 'Meta sin número en maestro; asignado por CUI')):
        ex = [p for p in ps if cond(p) and libre(p)]
        if ex:
            usados.add(id(ex[0])); return ex[0], nota
    ex = [p for p in ps if p['meta'] == meta]
    if ex:
        return ex[0], 'CUI/meta repetido en maestro; se usó la misma ubicación'
    return ps[0], 'Meta nueva respecto al maestro; se usó la ubicación del CUI'


# ---------------------------------------------------------------- v2: gasto corriente y sugerencias de Silucia
MAPA_EJECUTORAS_INICIAL = [   # ejecutora de Silucia -> gerencia/oficina (solo las que coincidieron >= 96% con la clasificación de Alex)
    ('GRI - POR ADM. DIRECTA', INFRA), ('GRI - POR CONTRATA', INFRA), ('GRI - ESTUDIOS DEFINITIVOS', INFRA), ('GRI - SGO', INFRA),
    ('GRI - MAQUINARIA', INFRA), ('GRI - CIPRL', INFRA), ('SGED', INFRA),
    ('GRDS - EJECUCION', 'GERENCIA REGIONAL DE DESARROLLO SOCIAL'), ('GRDS - ESTUDIOS DEFINITIVOS', 'GERENCIA REGIONAL DE DESARROLLO SOCIAL'),
    ('GRDE - EJECUCION', 'GERENCIA REGIONAL DE DESARROLLO ECONOMICO'), ('PROCOMPITE - EJECUCION', 'GERENCIA REGIONAL DE DESARROLLO ECONOMICO'),
    ('GRAA - EJECUCION', 'GERENCIA REGIONAL DE RECURSOS NATURALES Y GESTION AMBIENTAL'),
    ('PETT - EJECUCION', 'PROYECTO ESPECIAL TRUCHAS TITICACA - PETT'), ('PECSA - EJECUCION', 'PROYECTO ESPECIAL CAMELIDOS SUDAMERICANOS'),
    ('DEPI', 'DIRECCIÓN DE ESTUDIOS DE PREINVERSIÓN'),
]
HEAD_GC = ['UE', 'Meta', 'Gerencia / Oficina', 'Función', 'Actividad (referencia)', 'Origen']
HEAD_MAPA = ['Ejecutora en Silucia', 'Gerencia / Oficina sugerida']


def asegurar_estructura(path):
    """Agrega al maestro las hojas de la versión 2 si faltan (no toca lo existente)."""
    wb = openpyxl.load_workbook(path); cambio = False
    if 'ClasifGC' not in wb.sheetnames:
        ws = wb.create_sheet('ClasifGC'); _estilo_cabecera(ws, HEAD_GC)
        for col, w in zip('ABCDEF', [6, 8, 56, 34, 70, 22]):
            ws.column_dimensions[col].width = w
        cambio = True
    if 'MapaEjecutoras' not in wb.sheetnames:
        ws = wb.create_sheet('MapaEjecutoras'); _estilo_cabecera(ws, HEAD_MAPA)
        for r in MAPA_EJECUTORAS_INICIAL:
            ws.append(list(r))
        ws.column_dimensions['A'].width = 36; ws.column_dimensions['B'].width = 66
        cambio = True
    if 'Listas' in wb.sheetnames and 'ClasifGC' in wb.sheetnames:
        wl = wb['Listas']; n = len([c for c in wl['A'][1:] if c.value])
        if n and not wb['ClasifGC'].data_validations.dataValidation:
            dv = DataValidation(type='list', formula1=f"=Listas!$A$2:$A${n + 1}", allow_blank=True, showErrorMessage=False)
            wb['ClasifGC'].add_data_validation(dv); dv.add('C2:C2000'); cambio = True
    if cambio:
        wb.save(path)


def cargar_extras(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    gc = {}; mapa = {}
    if 'ClasifGC' in wb.sheetnames:
        for r in wb['ClasifGC'].iter_rows(min_row=2, values_only=True):
            if r[1]:
                gc[(str(r[0] or '001').zfill(3), _meta4(r[1]))] = (r[2] or '').strip()
    if 'MapaEjecutoras' in wb.sheetnames:
        for r in wb['MapaEjecutoras'].iter_rows(min_row=2, values_only=True):
            if r[0] and r[1]:
                mapa[str(r[0]).strip()] = str(r[1]).strip()
    return gc, mapa


def agregar_pendientes_gc(path, filas):
    """filas: [[UE, Meta, Gerencia/Oficina, Función, Actividad, Origen]] (Gerencia vacía = por completar)."""
    if not filas:
        return 0
    wb = openpyxl.load_workbook(path); ws = wb['ClasifGC']
    ex = {(str(r[0] or '001').zfill(3), _meta4(r[1])) for r in ws.iter_rows(min_row=2, values_only=True) if r[1]}
    n = 0
    for f in filas:
        if (f[0], f[1]) in ex:
            continue
        ws.append(f); n += 1
        ws.cell(ws.max_row, 3).fill = PatternFill('solid', fgColor='FFF2CC')
    wb.save(path)
    return n
