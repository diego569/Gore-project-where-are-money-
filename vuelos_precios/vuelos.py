# -*- coding: utf-8 -*-
"""Rastreador de precios de vuelos (Google Flights) Juliaca / Arequipa / Cusco <-> Lima.
Cada corrida agrega una "foto" de precios a historial/precios.csv y regenera vuelos_precios.xlsx.
"""
import csv, json, os, sys, time, datetime as dt
from collections import defaultdict

AQUI = os.path.dirname(os.path.abspath(__file__))
HIST = os.path.join(AQUI, "historial")
CSV_PRECIOS = os.path.join(HIST, "precios.csv")
CSV_MEJORES = os.path.join(HIST, "mejores.csv")
XLSX = os.path.join(AQUI, "vuelos_precios.xlsx")
CAMPOS = ["ts", "origen", "aeropuerto", "sentido", "fecha", "fecha_vuelta", "aerolinea", "sale", "llega",
          "paradas", "duracion_min", "precio", "moneda"]
CAMPOS_M = ["ts", "origen", "ida", "vuelta", "total_tarifa", "extras", "total_ajustado", "detalle", "tipo"]


def cargar_config():
    with open(os.path.join(AQUI, "config.json"), encoding="utf-8") as f:
        return json.load(f)


# ---------------------------------------------------------------- consulta a Google Flights
def _hm(v):
    v = (list(v or []) + [None, None])
    return "%02d:%02d" % (v[0] or 0, v[1] or 0)


def parsear(html):
    from selectolax.lexbor import LexborHTMLParser
    s = LexborHTMLParser(html).css_first(r"script.ds\:1")
    if s is None:
        raise RuntimeError("Google no devolvio resultados (posible pagina de consentimiento o bloqueo)")
    data = s.text().split("data:", 1)[1].rsplit(",", 1)[0]
    if data.endswith("errorHasStatus: true"):
        return []
    p = json.loads(data)
    out = []
    for bloque in (2, 3):  # 2 = otras opciones, 3 = mejores
        try:
            lista = p[bloque][0]
        except (IndexError, TypeError):
            lista = None
        for k in lista or []:
            try:
                fl = k[0]
                try:
                    precio = k[1][0][1]
                except (IndexError, TypeError):
                    precio = None
                segs = fl[2]
                out.append({
                    "aerolinea": "/".join(fl[1]),
                    "sale": _hm(segs[0][8]), "llega": _hm(segs[-1][10]),
                    "paradas": len(segs) - 1,
                    "duracion_min": sum((s[11] or 0) for s in segs),
                    "precio": precio,
                })
            except Exception:
                continue
    vistos, res = set(), []
    for r in out:
        key = (r["aerolinea"], r["sale"], r["llega"], r["paradas"])
        if key not in vistos:
            vistos.add(key)
            res.append(r)
    return res


def consultar(cfg, origen, destino, fecha, fecha_vuelta=None, intentos=3):
    from fast_flights import FlightQuery, Passengers, create_query, fetch_flights_html
    vuelos = [FlightQuery(date=fecha, from_airport=origen, to_airport=destino)]
    if fecha_vuelta:
        vuelos.append(FlightQuery(date=fecha_vuelta, from_airport=destino, to_airport=origen))
    q = create_query(flights=vuelos, seat=cfg["asientos"],
                     trip="round-trip" if fecha_vuelta else "one-way",
                     passengers=Passengers(adults=cfg["adultos"]), language="es", currency=cfg["moneda"])
    err = None
    for i in range(intentos):
        try:
            return parsear(fetch_flights_html(q))
        except Exception as e:  # reintenta con pausa creciente
            err = e
            time.sleep(4 * (i + 1))
    raise RuntimeError(str(err))


# ---------------------------------------------------------------- historial
def leer_csv(ruta):
    if not os.path.exists(ruta):
        return []
    with open(ruta, encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def anexar_csv(ruta, campos, filas):
    nuevo = not os.path.exists(ruta)
    with open(ruta, "a", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=campos)
        if nuevo:
            w.writeheader()
        w.writerows(filas)


# ---------------------------------------------------------------- logistica terrestre (Puno <-> AQP/CUZ)
def hora_a_float(hm):
    h, m = hm.split(":")
    return int(h) + int(m) / 60.0


def extras_ida(cfg, o, vuelo):
    """Costo extra y nota por ir de Puno al aeropuerto de origen (solo Arequipa/Cusco)."""
    if not o["terrestre"]:
        return 0, ""
    lg = cfg["logistica"]
    minimo = lg["hora_salida_puno"] + o["bus_horas"] + lg["colchon_aeropuerto_horas"]
    costo = o["bus_costo_por_tramo"]
    if hora_a_float(vuelo["sale"]) >= minimo:
        return costo, "Bus Puno->%s el mismo dia" % o["aeropuerto"]
    return costo + o["hotel_costo_noche"], "Salir de Puno el dia anterior (noche en %s)" % o["nombre"]


def extras_vuelta(cfg, o, vuelo):
    if not o["terrestre"]:
        return 0, ""
    lg = cfg["logistica"]
    llegada = hora_a_float(vuelo["llega"]) + lg["colchon_llegada_horas"] + o["bus_horas"]
    costo = o["bus_costo_por_tramo"]
    if llegada <= lg["hora_limite_llegar_a_puno"]:
        return costo, "Bus %s->Puno el mismo dia" % o["aeropuerto"]
    return costo + o["hotel_costo_noche"], "Dormir en %s y subir a Puno al dia siguiente" % o["nombre"]


def _desc(r):
    return "%s %s %s-%s %s" % (r["fecha"], r["aerolinea"], r["sale"], r["llega"],
                               "directo" if int(r["paradas"]) == 0 else "%s escala(s)" % r["paradas"])


def mejores_combinaciones(cfg, filas_ts):
    """Mejor ida+vuelta por origen sumando dos vuelos sueltos (con extras terrestres)
    y la mejor tarifa ida y vuelta consultada como viaje redondo."""
    res = []
    for o in cfg["origenes"]:
        ap = o["aeropuerto"]
        idas = [r for r in filas_ts if r["aeropuerto"] == ap and r["sentido"] == "ida" and r["precio"] not in ("", None)]
        vueltas = [r for r in filas_ts if r["aeropuerto"] == ap and r["sentido"] == "vuelta" and r["precio"] not in ("", None)]
        mejor = None
        for i in idas:
            ei, ni = extras_ida(cfg, o, i)
            for v in vueltas:
                if v["fecha"] <= i["fecha"]:
                    continue
                ev, nv = extras_vuelta(cfg, o, v)
                tarifa = float(i["precio"]) + float(v["precio"])
                total = tarifa + ei + ev
                if mejor is None or total < mejor[0]:
                    mejor = (total, tarifa, ei + ev, i, v, ni, nv)
        if mejor:
            total, tarifa, ex, i, v, ni, nv = mejor
            res.append({"origen": o["nombre"], "ida": _desc(i), "vuelta": _desc(v),
                        "total_tarifa": round(tarifa), "extras": round(ex), "total_ajustado": round(total),
                        "detalle": "; ".join(x for x in (ni, nv) if x), "tipo": "2 vuelos sueltos"})
        rt = [r for r in filas_ts if r["aeropuerto"] == ap and r["sentido"] == "ida_vuelta" and r["precio"] not in ("", None)]
        if rt:
            r = min(rt, key=lambda x: float(x["precio"]))
            ei, ni = extras_ida(cfg, o, r)
            ex = ei + (o["bus_costo_por_tramo"] if o["terrestre"] else 0)
            res.append({"origen": o["nombre"], "ida": _desc(r), "vuelta": r["fecha_vuelta"] + " (elegir en Google Flights)",
                        "total_tarifa": round(float(r["precio"])), "extras": round(ex),
                        "total_ajustado": round(float(r["precio"]) + ex), "detalle": ni,
                        "tipo": "ida y vuelta (tarifa conjunta)"})
    return res


# ---------------------------------------------------------------- Excel
def generar_excel(cfg, ts_actual):
    from openpyxl import Workbook
    from openpyxl.chart import LineChart, Reference
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter

    precios = leer_csv(CSV_PRECIOS)
    mejores = leer_csv(CSV_MEJORES)
    wb = Workbook()
    AZUL = PatternFill("solid", fgColor="1F4E78")
    VERDE = PatternFill("solid", fgColor="C6EFCE")

    def cab(ws, fila, textos):
        for j, t in enumerate(textos, 1):
            c = ws.cell(row=fila, column=j, value=t)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = AZUL
            c.alignment = Alignment(horizontal="center", wrap_text=True)

    def anchos(ws, w):
        for j, x in enumerate(w, 1):
            ws.column_dimensions[get_column_letter(j)].width = x

    # --- Resumen
    ws = wb.active
    ws.title = "Resumen"
    ws["A1"] = "Vuelos Puno (Juliaca / Arequipa / Cusco) <-> Lima, enero 2027"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = "Ultima consulta: %s   |   Moneda: %s   |   Fuente: Google Flights" % (ts_actual, cfg["moneda"])
    ws["A3"] = ("Extras = bus Puno<->aeropuerto y hotel cuando hace falta (estimados editables en config.json). "
                "Los precios de Google Flights cambian en el dia y pueden no incluir equipaje de bodega.")
    ws["A3"].alignment = Alignment(wrap_text=True)
    ws.merge_cells("A3:J3")
    ws.row_dimensions[3].height = 32
    cab(ws, 5, ["Origen", "Tipo", "Ida", "Vuelta", "Tarifa", "Extras", "TOTAL ajustado", "Minimo historico", "Vs. foto anterior", "Detalle logistico"])
    actual = [m for m in mejores if m["ts"] == ts_actual]
    previos = sorted({m["ts"] for m in mejores if m["ts"] < ts_actual})
    ts_prev = previos[-1] if previos else None
    fila = 6
    for m in sorted(actual, key=lambda x: float(x["total_ajustado"])):
        hist = [float(x["total_ajustado"]) for x in mejores if x["origen"] == m["origen"] and x["tipo"] == m["tipo"]]
        prev = [float(x["total_ajustado"]) for x in mejores if x["ts"] == ts_prev and x["origen"] == m["origen"] and x["tipo"] == m["tipo"]]
        delta = (float(m["total_ajustado"]) - prev[0]) if prev else None
        vals = [m["origen"], m["tipo"], m["ida"], m["vuelta"], float(m["total_tarifa"]), float(m["extras"]),
                float(m["total_ajustado"]), min(hist), delta if delta is not None else "primera foto", m["detalle"]]
        for j, v in enumerate(vals, 1):
            c = ws.cell(row=fila, column=j, value=v)
            if j in (5, 6, 7, 8) or (j == 9 and delta is not None):
                c.number_format = '#,##0;[Red]-#,##0'
        ws.cell(row=fila, column=7).font = Font(bold=True)
        if fila == 6:
            for j in range(1, 11):
                ws.cell(row=fila, column=j).fill = VERDE
        fila += 1
    anchos(ws, [11, 28, 42, 42, 10, 9, 15, 15, 15, 60])
    ws.freeze_panes = "A6"

    # --- Vuelos (foto actual)
    ws2 = wb.create_sheet("Vuelos actuales")
    cab(ws2, 1, ["Origen", "Sentido", "Fecha", "Fecha vuelta", "Aerolinea", "Sale", "Llega", "Paradas", "Duracion (min)", "Precio", "Nota logistica"])
    por_ap = {o["aeropuerto"]: o for o in cfg["origenes"]}
    filas_act = [r for r in precios if r["ts"] == ts_actual]
    filas_act.sort(key=lambda r: (r["origen"], r["sentido"], r["fecha"], float(r["precio"]) if r["precio"] else 1e9))
    for i, r in enumerate(filas_act, 2):
        o = por_ap[r["aeropuerto"]]
        nota = ""
        if r["sentido"] == "ida":
            nota = extras_ida(cfg, o, r)[1]
        elif r["sentido"] == "vuelta":
            nota = extras_vuelta(cfg, o, r)[1]
        vals = [r["origen"], r["sentido"], r["fecha"], r["fecha_vuelta"], r["aerolinea"], r["sale"], r["llega"],
                int(r["paradas"]), int(r["duracion_min"] or 0), float(r["precio"]) if r["precio"] else "sin precio", nota]
        for j, v in enumerate(vals, 1):
            ws2.cell(row=i, column=j, value=v)
        ws2.cell(row=i, column=10).number_format = "#,##0"
    anchos(ws2, [11, 11, 12, 12, 18, 8, 8, 9, 13, 11, 52])
    ws2.freeze_panes = "A2"
    ws2.auto_filter.ref = "A1:K%d" % max(2, len(filas_act) + 1)

    # --- Historial: precio minimo por origen / sentido / fecha en cada foto
    ws3 = wb.create_sheet("Historial")
    claves = sorted({(r["origen"], r["sentido"], r["fecha"], r["fecha_vuelta"]) for r in precios})
    claves = [k for k in claves if k[1] in ("ida", "vuelta")]
    tss = sorted({r["ts"] for r in precios})
    cab(ws3, 1, ["Foto (fecha-hora)"] + ["%s %s %s" % (k[0], k[1], k[2][5:]) for k in claves])
    mins = {}
    for r in precios:
        if r["precio"]:
            k = (r["ts"], r["origen"], r["sentido"], r["fecha"], r["fecha_vuelta"])
            p = float(r["precio"])
            mins[k] = p if k not in mins else min(mins[k], p)
    for i, t in enumerate(tss, 2):
        ws3.cell(row=i, column=1, value=t)
        for j, k in enumerate(claves, 2):
            ws3.cell(row=i, column=j, value=mins.get((t,) + k))
    anchos(ws3, [18] + [15] * len(claves))
    ws3.freeze_panes = "B2"

    # --- Evolucion del mejor total por origen (grafico)
    ws4 = wb.create_sheet("Evolucion")
    origs = [o["nombre"] for o in cfg["origenes"]]
    cab(ws4, 1, ["Foto"] + origs)
    ts_m = sorted({m["ts"] for m in mejores})
    for i, t in enumerate(ts_m, 2):
        ws4.cell(row=i, column=1, value=t)
        for j, nom in enumerate(origs, 2):
            c = [float(m["total_ajustado"]) for m in mejores if m["ts"] == t and m["origen"] == nom and m["tipo"] == "2 vuelos sueltos"]
            ws4.cell(row=i, column=j, value=min(c) if c else None)
    anchos(ws4, [18, 14, 14, 14])
    if ts_m:
        ch = LineChart()
        ch.title = "Mejor total ajustado (2 vuelos sueltos + extras) por origen"
        ch.y_axis.title = cfg["moneda"]
        ch.height, ch.width = 9, 22
        ch.add_data(Reference(ws4, min_col=2, max_col=1 + len(origs), min_row=1, max_row=1 + len(ts_m)), titles_from_data=True)
        ch.set_categories(Reference(ws4, min_col=1, min_row=2, max_row=1 + len(ts_m)))
        ws4.add_chart(ch, "F2")
    wb.save(XLSX)


# ---------------------------------------------------------------- principal
def main():
    cfg = cargar_config()
    os.makedirs(HIST, exist_ok=True)
    ts = dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    destino = cfg["destino"]
    filas, errores = [], []

    tareas = []
    for o in cfg["origenes"]:
        ap = o["aeropuerto"]
        for f in cfg["fechas_ida"]:
            tareas.append((o, "ida", ap, destino, f, None))
        for f in cfg["fechas_vuelta"]:
            tareas.append((o, "vuelta", destino, ap, f, None))
        if cfg.get("consultar_ida_y_vuelta"):
            for f in cfg["fechas_ida"]:
                for g in cfg["fechas_vuelta"]:
                    tareas.append((o, "ida_vuelta", ap, destino, f, g))

    for n, (o, sentido, a, b, f, g) in enumerate(tareas, 1):
        etiqueta = "%s %s->%s %s%s" % (sentido, a, b, f, (" / " + g) if g else "")
        try:
            res = consultar(cfg, a, b, f, g)
        except Exception as e:
            errores.append("%s: %s" % (etiqueta, e))
            print("[%d/%d] %s  ERROR" % (n, len(tareas), etiqueta))
            continue
        for r in res:
            filas.append({"ts": ts, "origen": o["nombre"], "aeropuerto": o["aeropuerto"], "sentido": sentido, "fecha": f,
                          "fecha_vuelta": g or "", "aerolinea": r["aerolinea"], "sale": r["sale"], "llega": r["llega"],
                          "paradas": r["paradas"], "duracion_min": r["duracion_min"],
                          "precio": r["precio"] if r["precio"] is not None else "", "moneda": cfg["moneda"]})
        con = [r["precio"] for r in res if r["precio"]]
        print("[%d/%d] %s  %d vuelos, desde %s" % (n, len(tareas), etiqueta, len(res), min(con) if con else "-"))
        time.sleep(cfg.get("pausa_segundos", 2.5))

    if not filas:
        print("No se obtuvo ningun resultado. Errores:\n" + "\n".join(errores))
        sys.exit(1)

    anexar_csv(CSV_PRECIOS, CAMPOS, filas)
    mej = mejores_combinaciones(cfg, filas)
    anexar_csv(CSV_MEJORES, CAMPOS_M, [dict(m, ts=ts) for m in mej])
    generar_excel(cfg, ts)

    todos = leer_csv(CSV_MEJORES)
    previos = sorted({m["ts"] for m in todos if m["ts"] < ts})
    ts_prev = previos[-1] if previos else None
    lineas = ["Consulta %s" % ts]
    for m in sorted(mej, key=lambda x: x["total_ajustado"]):
        prev = [float(x["total_ajustado"]) for x in todos if x["ts"] == ts_prev and x["origen"] == m["origen"] and x["tipo"] == m["tipo"]]
        cambio = ""
        if prev:
            d = m["total_ajustado"] - prev[0]
            cambio = "  (%+d vs foto anterior)" % d
            if d < 0 and abs(d) / prev[0] * 100 >= cfg.get("alerta_baja_porcentaje", 5):
                cambio += "  <<< BAJO"
        lineas.append("%-9s %-30s total %s %d (tarifa %d + extras %d)%s\n          ida %s | vuelta %s" % (
            m["origen"], m["tipo"], cfg["moneda"], m["total_ajustado"], m["total_tarifa"], m["extras"], cambio, m["ida"], m["vuelta"]))
    if errores:
        lineas.append("Consultas con error (%d): %s" % (len(errores), " | ".join(errores)))
    texto = "\n".join(lineas)
    with open(os.path.join(AQUI, "ULTIMO_RESULTADO.txt"), "w", encoding="utf-8") as f:
        f.write(texto)
    print("\n" + texto)


if __name__ == "__main__":
    main()
